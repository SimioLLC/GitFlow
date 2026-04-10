"""Convert the v2 user guide markdown to a formatted .docx file using python-docx."""

import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

INPUT = "Simio with Version Control Software Overview v2.md"
OUTPUT = "Simio with Version Control Software Overview v2.docx"


def set_cell_shading(cell, color_hex):
    """Set background shading on a table cell."""
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn("w:shd"), {
        qn("w:val"): "clear",
        qn("w:color"): "auto",
        qn("w:fill"): color_hex,
    })
    shading.append(shd)


def style_document(doc):
    """Configure document-level styles."""
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)

    for level, size, bold in [
        ("Heading 1", 22, True),
        ("Heading 2", 16, True),
        ("Heading 3", 13, True),
        ("Heading 4", 11, True),
    ]:
        s = doc.styles[level]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.bold = bold
        s.font.color.rgb = RGBColor(0x1A, 0x47, 0x8A)
        s.paragraph_format.space_before = Pt(12)
        s.paragraph_format.space_after = Pt(4)


def add_formatted_paragraph(doc, text, style="Normal", bold=False):
    """Add a paragraph with inline bold (**text**) and italic (*text*) support."""
    p = doc.add_paragraph(style=style)
    # Split on bold and italic markers
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and not part.startswith("**"):
            run = p.add_run(part[1:-1])
            run.italic = True
        else:
            p.add_run(part)
    if bold:
        for run in p.runs:
            run.bold = True
    return p


def add_table(doc, header_row, data_rows):
    """Add a formatted table."""
    cols = len(header_row)
    table = doc.add_table(rows=1 + len(data_rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    for i, text in enumerate(header_row):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(cell, "1A478A")

    # Data rows
    for r, row_data in enumerate(data_rows):
        for c, text in enumerate(row_data):
            cell = table.rows[r + 1].cells[c]
            cell.text = ""
            p = cell.paragraphs[0]
            # Support bold in table cells
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                    run.font.size = Pt(10)
                else:
                    run = p.add_run(part)
                    run.font.size = Pt(10)
            # Alternate row shading
            if r % 2 == 1:
                set_cell_shading(cell, "E8EDF5")

    return table


def parse_table(lines, start_idx):
    """Parse a markdown table starting at the given line index. Returns (header, rows, end_idx)."""
    header_line = lines[start_idx].strip()
    header = [c.strip().strip("*") for c in header_line.split("|") if c.strip()]
    # Skip separator line
    data_rows = []
    idx = start_idx + 2  # skip header + separator
    while idx < len(lines) and "|" in lines[idx] and not lines[idx].strip().startswith("#"):
        row = [c.strip() for c in lines[idx].split("|") if c.strip()]
        if row:
            data_rows.append(row)
        idx += 1
    return header, data_rows, idx


def convert():
    with open(INPUT, "r", encoding="utf-8") as f:
        content = f.read()

    lines = content.split("\n")
    doc = Document()
    style_document(doc)

    # Set narrow margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    i = 0
    in_list = False
    list_indent = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Skip horizontal rules
        if stripped == "---":
            i += 1
            continue

        # Headings
        if stripped.startswith("####"):
            text = stripped.lstrip("#").strip()
            doc.add_heading(text, level=4)
            i += 1
            continue
        if stripped.startswith("###"):
            text = stripped.lstrip("#").strip()
            doc.add_heading(text, level=3)
            i += 1
            continue
        if stripped.startswith("##"):
            text = stripped.lstrip("#").strip()
            doc.add_heading(text, level=2)
            i += 1
            continue
        if stripped.startswith("#"):
            text = stripped.lstrip("#").strip()
            doc.add_heading(text, level=1)
            i += 1
            continue

        # Tables
        if "|" in stripped and i + 1 < len(lines) and re.match(r'\|[\s\-:|]+\|', lines[i + 1].strip()):
            header, data_rows, end_idx = parse_table(lines, i)
            add_table(doc, header, data_rows)
            doc.add_paragraph()  # spacing after table
            i = end_idx
            continue

        # Numbered lists
        num_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if num_match:
            text = num_match.group(2)
            add_formatted_paragraph(doc, text, style="List Number")
            i += 1
            continue

        # Bullet lists (- or *)
        bullet_match = re.match(r'^[-*]\s+(.*)', stripped)
        if bullet_match:
            text = bullet_match.group(1)
            add_formatted_paragraph(doc, text, style="List Bullet")
            i += 1
            continue

        # Indented bullets (sub-items)
        sub_bullet_match = re.match(r'^\s+[-*]\s+(.*)', line)
        if sub_bullet_match:
            text = sub_bullet_match.group(1)
            p = add_formatted_paragraph(doc, text, style="List Bullet 2")
            i += 1
            continue

        # Empty lines
        if not stripped:
            i += 1
            continue

        # Regular paragraphs - collect continuation lines
        para_text = stripped
        i += 1
        while i < len(lines):
            next_stripped = lines[i].strip()
            if (not next_stripped or next_stripped.startswith("#") or
                next_stripped.startswith("-") or next_stripped.startswith("*") or
                re.match(r'^\d+\.', next_stripped) or "|" in next_stripped):
                break
            para_text += " " + next_stripped
            i += 1

        add_formatted_paragraph(doc, para_text)

    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    convert()
